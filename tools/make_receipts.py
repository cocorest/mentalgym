"""イルカマスターコースの領収書PDFを受講生名簿(xlsx)から作成する。

使い方: python3 tools/make_receipts.py 名簿.xlsx 出力フォルダ
"""
import sys
from datetime import date
from pathlib import Path

import openpyxl
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ISSUER_LINES = ["臨床心理士・公認心理師", "ぽ子"]
PURPOSE = "イルカマスターコース受講料として"
ISSUE_DATE = date(2026, 10, 10)
DEFAULT_AMOUNT = 150_000
AMOUNT_BY_NICKNAME = {"ちくわちゃん": 200_000}

# 入金済み・領収書希望○の人だけ作る。宛名が空欄なら本名を使う。
PAID = {"〇", "○"}

pdfmetrics.registerFont(TTFont("IPAG", "/usr/share/fonts/opentype/ipafont-gothic/ipag.ttf"))
FONT = "IPAG"
ROSE = (0.77, 0.41, 0.48)
BROWN = (0.29, 0.19, 0.21)


def load_rows(xlsx):
    ws = openpyxl.load_workbook(xlsx, data_only=True)["マスターコース"]
    header = [c.value for c in ws[3]]
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not isinstance(row[0], int):
            break
        yield dict(zip(header, row))


def draw_receipt(path, no, addressee, amount):
    w, h = landscape(A4)
    c = canvas.Canvas(str(path), pagesize=(w, h))
    c.setTitle(f"領収書 {addressee}")
    c.setStrokeColorRGB(*ROSE)
    c.setLineWidth(2)
    c.rect(15 * mm, 15 * mm, w - 30 * mm, h - 30 * mm)
    c.setLineWidth(0.5)
    c.rect(17 * mm, 17 * mm, w - 34 * mm, h - 34 * mm)

    c.setFillColorRGB(*BROWN)
    c.setFont(FONT, 30)
    c.drawCentredString(w / 2, h - 45 * mm, "領　収　書")
    c.setFont(FONT, 10)
    c.drawRightString(w - 25 * mm, h - 28 * mm, f"No. {no}")
    c.drawRightString(w - 25 * mm, h - 34 * mm, f"発行日 {ISSUE_DATE.year}年{ISSUE_DATE.month}月{ISSUE_DATE.day}日")

    # 宛名
    c.setFont(FONT, 18)
    honorific = "御中" if any(k in addressee for k in ("株式会社", "（株）", "有限会社", "合同会社")) else "様"
    name = f"{addressee}　{honorific}"
    c.drawString(35 * mm, h - 70 * mm, name)
    c.setStrokeColorRGB(*BROWN)
    c.line(35 * mm, h - 72 * mm, max(150 * mm, 35 * mm + c.stringWidth(name, FONT, 18) + 10 * mm), h - 72 * mm)

    # 金額
    box_w, box_h = 150 * mm, 22 * mm
    bx, by = (w - box_w) / 2, h - 108 * mm
    c.setFillColorRGB(0.99, 0.94, 0.95)
    c.rect(bx, by, box_w, box_h, fill=1, stroke=0)
    c.setFillColorRGB(*BROWN)
    c.setFont(FONT, 26)
    c.drawCentredString(w / 2, by + 7 * mm, f"¥{amount:,}－")

    c.setFont(FONT, 12)
    c.drawString(bx, by - 10 * mm, f"但し　{PURPOSE}")
    c.drawString(bx, by - 17 * mm, "上記正に領収いたしました。")

    # 収入印紙欄
    c.setDash(2, 2)
    c.rect(35 * mm, 28 * mm, 22 * mm, 26 * mm)
    c.setDash()
    c.setFont(FONT, 8)
    c.drawCentredString(46 * mm, 40 * mm, "収入印紙")

    # 発行者
    c.setFont(FONT, 13)
    y = 50 * mm
    for line in ISSUER_LINES:
        c.drawRightString(w - 35 * mm, y, line)
        y -= 8 * mm
    c.showPage()
    c.save()


def main(xlsx, out_dir):
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    made, held = [], []
    for r in load_rows(xlsx):
        nick = r["受講生名"]
        if r["領収書希望"] not in PAID:
            continue
        if r["入金状況"] not in PAID:
            held.append((nick, r["入金状況"]))
            continue
        addressee = (r["領収書宛名"] or r["本名"] or "").strip()
        no = f"IM-{ISSUE_DATE.year}-{r['No.']:03d}"
        amount = AMOUNT_BY_NICKNAME.get(nick, DEFAULT_AMOUNT)
        path = out / f"領収書_{r['No.']:02d}_{nick}_{addressee.replace(' ', '').replace('　', '')}.pdf"
        draw_receipt(path, no, addressee, amount)
        made.append((no, nick, addressee, amount, "宛名欄空欄→本名" if not r["領収書宛名"] else ""))
    for m in made:
        print("作成", *m)
    for h in held:
        print("保留", *h)


if __name__ == "__main__":
    main(*sys.argv[1:3])
